from django.shortcuts import render, redirect, get_object_or_404
from django.views.generic import TemplateView, ListView, UpdateView, CreateView, DeleteView
from django.contrib.auth.mixins import LoginRequiredMixin
from django.contrib.auth.decorators import login_required
from django.urls import reverse_lazy
from django.http import HttpResponse, JsonResponse
from django import forms
from django.contrib import messages
from apps.core.models import Event, Guest, Table, Photo, ChecklistItem, BudgetItem
from apps.core.views import PlannerRequiredMixin
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def get_user_event(user):
    if user.is_authenticated and hasattr(user, 'event') and user.event:
        return user.event
    return Event.objects.first()


class DashboardHomeView(PlannerRequiredMixin, TemplateView):
    template_name = 'dashboard/dashboard.html'
    
    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        event = get_user_event(self.request.user)
        
        if event:
            main_guests = event.guests.filter(linked_to__isnull=True)
            context['event'] = event
            context['days_left'] = event.days_until
            context['total_guests'] = main_guests.count()
            context['confirmed_guests'] = main_guests.filter(rsvp_status='CONFIRMED').count()
            context['pending_guests'] = context['total_guests'] - context['confirmed_guests']
            
            checklist_items = event.checklist_items.all()
            context['completed_tasks'] = checklist_items.filter(is_completed=True).count()
            context['total_tasks'] = checklist_items.count()
            context['checklist_items'] = checklist_items[:4]
            
            budget_items = event.budget_items.all()
            total_estimated = sum(item.estimated_cost for item in budget_items)
            total_actual = sum(item.actual_cost for item in budget_items)
            context['total_estimated'] = total_estimated
            context['total_actual'] = total_actual
            context['budget_total'] = event.budget_total
            context['budget_items'] = budget_items[:5]
            
            context['recent_photos'] = event.photos.all()[:6]
            context['unassigned_guests'] = main_guests.filter(assigned_table__isnull=True).count()
            context['tables'] = event.tables.all()
        else:
            context['event'] = None
            context['days_left'] = 0
            context['total_guests'] = 0
            context['confirmed_guests'] = 0
            context['pending_guests'] = 0
            context['completed_tasks'] = 0
            context['total_tasks'] = 0
            context['checklist_items'] = []
            context['total_estimated'] = 0
            context['total_actual'] = 0
            context['budget_total'] = 0
            context['budget_items'] = []
            context['recent_photos'] = []
            context['unassigned_guests'] = 0
            context['tables'] = []
        
        return context


class ChecklistView(PlannerRequiredMixin, ListView):
    template_name = 'dashboard/checklist.html'
    model = ChecklistItem
    context_object_name = 'checklist_items'
    
    def get_queryset(self):
        event = get_user_event(self.request.user)
        if event:
            return event.checklist_items.all()
        return ChecklistItem.objects.none()
    
    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['event'] = get_user_event(self.request.user)
        return context


@login_required
def toggle_checklist_item(request, pk):
    if request.user.role != 'PLANNER':
        return redirect('dashboard:home')
    
    item = get_object_or_404(ChecklistItem, pk=pk)
    item.is_completed = not item.is_completed
    item.save()
    return redirect('dashboard:checklist')


class ChecklistItemForm(forms.ModelForm):
    class Meta:
        model = ChecklistItem
        fields = ['title', 'priority', 'due_date']
        widgets = {
            'title': forms.TextInput(attrs={'placeholder': 'Descripción de la tarea'}),
            'due_date': forms.DateInput(attrs={'type': 'date'}),
        }


class ChecklistItemCreateView(PlannerRequiredMixin, CreateView):
    model = ChecklistItem
    form_class = ChecklistItemForm
    template_name = 'dashboard/checklist_form.html'
    
    def get_success_url(self):
        return reverse_lazy('dashboard:checklist')
    
    def form_valid(self, form):
        event = get_user_event(self.request.user)
        if not event:
            messages.error(self.request, 'Primero debes configurar tu evento.')
            return redirect('dashboard:event_settings')
        form.instance.event = event
        messages.success(self.request, f'Tarea "{form.instance.title}" agregada.')
        return super().form_valid(form)


class ChecklistItemUpdateView(PlannerRequiredMixin, UpdateView):
    model = ChecklistItem
    form_class = ChecklistItemForm
    template_name = 'dashboard/checklist_form.html'
    
    def get_success_url(self):
        return reverse_lazy('dashboard:checklist')
    
    def form_valid(self, form):
        messages.success(self.request, f'Tarea "{form.instance.title}" actualizada.')
        return super().form_valid(form)


class ChecklistItemDeleteView(PlannerRequiredMixin, DeleteView):
    model = ChecklistItem
    template_name = 'dashboard/checklist_confirm_delete.html'
    
    def get_success_url(self):
        return reverse_lazy('dashboard:checklist')
    
    def delete(self, request, *args, **kwargs):
        messages.success(request, 'Tarea eliminada.')
        return super().delete(request, *args, **kwargs)


class EventForm(forms.ModelForm):
    DESIGN_PREVIEW_MAP = {
        'classic_romance': {
            'name': 'Romance Clásico',
            'description': 'Elegante y romántico con tonos rosa y dorado. Tipografía serif con acentos ornamentales.',
            'colors': ['#DB2777', '#F472B6', '#CA8A04', '#FDF2F8'],
            'fonts': 'Great Vibes + Cormorant Infant',
        },
        'modern_minimal': {
            'name': 'Minimalista Moderno',
            'description': 'Limpio y sofisticado con alto contraste. Diseño contemporáneo con espacios amplios.',
            'colors': ['#1A1A1A', '#FFFFFF', '#CA8A04', '#F9FAFB'],
            'fonts': 'Bodoni Moda + Jost',
        },
        'rustic_boho': {
            'name': 'Rústico Boho',
            'description': 'Cálido y natural con tonos tierra y verde sage. Elementos botánicos y textura orgánica.',
            'colors': ['#92400E', '#B45309', '#6B8E6B', '#FEF3C7'],
            'fonts': 'Cormorant + Montserrat',
        },
        'art_deco': {
            'name': 'Art Deco',
            'description': 'Glamouroso y vintage con patrones geométricos dorados sobre fondo oscuro. Estilo 1920s.',
            'colors': ['#CA8A04', '#0A0A0A', '#F5F0E6', '#1A1A1A'],
            'fonts': 'Poiret One + Didact Gothic',
        },
        'floral_romantic': {
            'name': 'Floral Romántico',
            'description': 'Delicado y femenino con acuarelas florales. Rosa suave y verde sage con tipografía script.',
            'colors': ['#E8A0BF', '#DB2777', '#A7C4A0', '#FFF8F0'],
            'fonts': 'Great Vibes + Cormorant Infant',
        },
    }
    
    class Meta:
        model = Event
        fields = ['bride_name', 'groom_name', 'date', 'location', 'venue_details', 'latitude', 'longitude', 'couple_photo', 'design_template', 'email_header', 'email_message', 'email_question', 'email_button_text', 'email_closing']
        widgets = {
            'bride_name': forms.TextInput(attrs={'placeholder': 'Nombre de la novia'}),
            'groom_name': forms.TextInput(attrs={'placeholder': 'Nombre del novio'}),
            'date': forms.DateInput(attrs={'type': 'date'}),
            'location': forms.TextInput(attrs={'placeholder': 'Nombre del lugar (ej: Hacienda Los Nogales)'}),
            'venue_details': forms.Textarea(attrs={'placeholder': 'Direccion completa, indicaciones para los invitados...', 'rows': 3}),
            'latitude': forms.HiddenInput(),
            'longitude': forms.HiddenInput(),
            'design_template': forms.RadioSelect(),
            'email_header': forms.TextInput(attrs={'placeholder': 'Estás invitado'}),
            'email_message': forms.Textarea(attrs={'placeholder': 'Nos llena de alegría compartir este momento tan especial contigo', 'rows': 3}),
            'email_question': forms.TextInput(attrs={'placeholder': '¿nos acompañarás?'}),
            'email_button_text': forms.TextInput(attrs={'placeholder': 'Confirmar Asistencia'}),
            'email_closing': forms.TextInput(attrs={'placeholder': 'Con amor'}),
        }


class EventSettingsView(PlannerRequiredMixin, TemplateView):
    template_name = 'dashboard/event_settings.html'
    
    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        event = get_user_event(self.request.user)
        if event:
            context['form'] = EventForm(instance=event)
            context['current_lat'] = event.latitude or -33.4569
            context['current_lng'] = event.longitude or -70.6483
        else:
            context['form'] = EventForm()
            context['current_lat'] = -33.4569
            context['current_lng'] = -70.6483
        context['event'] = event
        return context
    
    def post(self, request, *args, **kwargs):
        event = get_user_event(request.user)
        if event:
            form = EventForm(request.POST, request.FILES, instance=event)
        else:
            form = EventForm(request.POST, request.FILES)
        
        if form.is_valid():
            event = form.save()
            if not request.user.event:
                request.user.event = event
                request.user.save()
            return redirect('dashboard:home')
        
        context = self.get_context_data(**kwargs)
        context['form'] = form
        return render(request, self.template_name, context)


class GuestForm(forms.ModelForm):
    class Meta:
        model = Guest
        fields = ['full_name', 'email', 'phone', 'group', 'rsvp_status', 'meal_choice', 'plus_one', 'plus_one_name', 'notes']
        widgets = {
            'full_name': forms.TextInput(attrs={'placeholder': 'Nombre completo'}),
            'email': forms.EmailInput(attrs={'placeholder': 'correo@ejemplo.com'}),
            'phone': forms.TextInput(attrs={'placeholder': 'Teléfono'}),
            'notes': forms.Textarea(attrs={'placeholder': 'Notas o alergias...', 'rows': 2}),
            'plus_one_name': forms.TextInput(attrs={'placeholder': 'Nombre del acompañante'}),
        }
    
    def __init__(self, *args, **kwargs):
        self.event = kwargs.pop('event', None)
        super().__init__(*args, **kwargs)
    
    def clean_email(self):
        email = self.cleaned_data.get('email')
        if email and self.event:
            queryset = Guest.objects.filter(event=self.event, email=email)
            if self.instance.pk:
                queryset = queryset.exclude(pk=self.instance.pk)
            if queryset.exists():
                raise forms.ValidationError('Ya existe un invitado con este correo electrónico en este evento.')
        return email


class GuestListView(PlannerRequiredMixin, ListView):
    template_name = 'dashboard/guest_list.html'
    model = Guest
    context_object_name = 'guests'
    
    def get_queryset(self):
        event = get_user_event(self.request.user)
        if not event:
            return Guest.objects.none()
        
        queryset = event.guests.filter(linked_to__isnull=True)
        
        search = self.request.GET.get('search', '')
        if search:
            queryset = queryset.filter(full_name__icontains=search) | queryset.filter(email__icontains=search)
        
        group_filter = self.request.GET.get('group', '')
        if group_filter:
            queryset = queryset.filter(group=group_filter)
        
        rsvp_filter = self.request.GET.get('rsvp', '')
        if rsvp_filter:
            queryset = queryset.filter(rsvp_status=rsvp_filter)
        
        return queryset.distinct()
    
    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['event'] = get_user_event(self.request.user)
        context['search'] = self.request.GET.get('search', '')
        context['group_filter'] = self.request.GET.get('group', '')
        context['rsvp_filter'] = self.request.GET.get('rsvp', '')
        if context['event']:
            main_guests = context['event'].guests.filter(linked_to__isnull=True)
            context['total_guests'] = main_guests.count()
            context['confirmed_guests'] = main_guests.filter(rsvp_status='CONFIRMED').count()
        else:
            context['total_guests'] = 0
            context['confirmed_guests'] = 0
        context['pending_guests'] = context['total_guests'] - context['confirmed_guests']
        context['tables'] = context['event'].tables.all() if context['event'] else []
        return context


class GuestCreateView(PlannerRequiredMixin, CreateView):
    model = Guest
    form_class = GuestForm
    template_name = 'dashboard/guest_form.html'
    
    def get_success_url(self):
        return reverse_lazy('dashboard:guests')
    
    def get_form_kwargs(self):
        kwargs = super().get_form_kwargs()
        kwargs['event'] = get_user_event(self.request.user)
        return kwargs
    
    def form_valid(self, form):
        event = get_user_event(self.request.user)
        if not event:
            messages.error(self.request, 'Primero debes configurar tu evento.')
            return redirect('dashboard:event_settings')
        form.instance.event = event
        
        existing_guest = Guest.objects.filter(event=event, email=form.cleaned_data['email']).first()
        if existing_guest and existing_guest.invitation_sent:
            messages.warning(self.request, f'El invitado {form.instance.full_name} ya fue agregado y el correo de invitación ya fue enviado anteriormente.')
            form.instance = existing_guest
            return redirect(self.get_success_url())
        
        response = super().form_valid(form)
        
        try:
            from apps.core.emails import send_invitation_email
            send_invitation_email(self.object)
            messages.success(self.request, f'Invitado {form.instance.full_name} agregado y correo de invitación enviado.')
        except Exception as e:
            messages.warning(self.request, f'Invitado agregado, pero no se pudo enviar el correo: {str(e)}')
        
        return response


class GuestUpdateView(PlannerRequiredMixin, UpdateView):
    model = Guest
    form_class = GuestForm
    template_name = 'dashboard/guest_form.html'
    
    def get_success_url(self):
        return reverse_lazy('dashboard:guests')
    
    def get_form_kwargs(self):
        kwargs = super().get_form_kwargs()
        kwargs['event'] = get_user_event(self.request.user)
        return kwargs
    
    def form_valid(self, form):
        messages.success(self.request, f'Invitado {form.instance.full_name} actualizado correctamente.')
        return super().form_valid(form)


class GuestDeleteView(PlannerRequiredMixin, DeleteView):
    model = Guest
    template_name = 'dashboard/guest_confirm_delete.html'
    
    def get_success_url(self):
        return reverse_lazy('dashboard:guests')
    
    def delete(self, request, *args, **kwargs):
        messages.success(request, 'Invitado eliminado correctamente.')
        return super().delete(request, *args, **kwargs)


@login_required
def import_guests(request):
    if request.user.role != 'PLANNER':
        return JsonResponse({'error': 'No autorizado'}, status=403)
    
    event = get_user_event(request.user)
    if not event:
        return JsonResponse({'error': 'No hay evento configurado'}, status=400)
    
    if request.method != 'POST':
        return JsonResponse({'error': 'Método no permitido'}, status=405)
    
    file = request.FILES.get('file')
    if not file:
        return JsonResponse({'error': 'No se proporcionó archivo'}, status=400)
    
    try:
        from openpyxl import load_workbook
        from apps.core.emails import send_invitation_email
        
        wb = load_workbook(file)
        ws = wb.active
        
        headers = [cell.value for cell in ws[1]]
        required_cols = ['nombre', 'email', 'grupo']
        
        header_map = {}
        for i, header in enumerate(headers):
            if header:
                header_map[header.lower().strip()] = i
        
        for col in required_cols:
            if col not in header_map:
                return JsonResponse({'error': f'Columna requerida faltante: {col}'}, status=400)
        
        created_count = 0
        updated_count = 0
        emails_sent = 0
        errors = []
        
        for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            try:
                name = row[header_map['nombre']] if 'nombre' in header_map else None
                email = row[header_map['email']] if 'email' in header_map else None
                group = row[header_map['grupo']].upper().strip() if 'grupo' in header_map and row[header_map['grupo']] else None
                phone = row[header_map['telefono']].strip() if 'telefono' in header_map and row[header_map['telefono']] else ''
                meal = row[header_map['menu']].upper().strip() if 'menu' in header_map and row[header_map['menu']] else 'STANDARD'
                plus_one = bool(row[header_map['acompañante']]) if 'acompañante' in header_map and row[header_map['acompañante']] else False
                plus_one_name = str(row[header_map['acompañante']]) if 'acompañante' in header_map and row[header_map['acompañante']] and not isinstance(row[header_map['acompañante']], bool) else ''
                notes = row[header_map['notas']] if 'notas' in header_map else ''
                
                if not name or not email:
                    errors.append(f'Fila {row_num}: nombre y email son requeridos')
                    continue
                
                group_map = {
                    'FAMILIA NOVIA': 'NOVIA',
                    'NOVIA': 'NOVIA',
                    'FAMILIA NOVIO': 'NOVIO',
                    'NOVIO': 'NOVIO',
                    'AMIGOS NOVIA': 'AMIGOS_NOVIA',
                    'AMIGOS DE LA NOVIA': 'AMIGOS_NOVIA',
                    'AMIGOS NOVIO': 'AMIGOS_NOVIO',
                    'AMIGOS DEL NOVIO': 'AMIGOS_NOVIO',
                    'TRABAJO': 'TRABAJO',
                    'COMPANEROS': 'TRABAJO',
                    'OTROS': 'OTROS',
                }
                
                if group and group in group_map:
                    group = group_map[group]
                elif group and group not in [c[0] for c in Guest.GROUP_CHOICES]:
                    errors.append(f'Fila {row_num}: grupo inválido "{group}"')
                    continue
                
                if meal not in [c[0] for c in Guest.MEAL_CHOICES]:
                    meal = 'STANDARD'
                
                existing_guest = Guest.objects.filter(event=event, email=email).first()
                
                if existing_guest and existing_guest.invitation_sent:
                    updated_count += 1
                    continue
                
                guest, created = Guest.objects.update_or_create(
                    event=event,
                    email=email,
                    defaults={
                        'full_name': name,
                        'phone': phone,
                        'group': group or 'OTROS',
                        'meal_choice': meal,
                        'plus_one': plus_one,
                        'plus_one_name': plus_one_name,
                        'notes': notes or '',
                    }
                )
                
                if created:
                    created_count += 1
                    try:
                        send_invitation_email(guest)
                        emails_sent += 1
                    except Exception as e:
                        errors.append(f'Fila {row_num}: Error al enviar correo a {email}: {str(e)}')
                else:
                    updated_count += 1
                    if not guest.invitation_sent:
                        try:
                            send_invitation_email(guest)
                            emails_sent += 1
                        except Exception as e:
                            errors.append(f'Fila {row_num}: Error al enviar correo a {email}: {str(e)}')
                    
            except Exception as e:
                errors.append(f'Fila {row_num}: {str(e)}')
        
        message = f'Importación completada: {created_count} nuevos, {updated_count} actualizados.'
        if emails_sent > 0:
            message += f' {emails_sent} correos enviados.'
        if errors:
            message += f' {len(errors)} errores.'
        
        return JsonResponse({
            'success': True,
            'created': created_count,
            'updated': updated_count,
            'emails_sent': emails_sent,
            'errors': errors[:20]
        })
        
    except Exception as e:
        return JsonResponse({'error': f'Error al procesar archivo: {str(e)}'}, status=400)


@login_required
def export_guests(request):
    if request.user.role != 'PLANNER':
        return HttpResponse('No autorizado', status=403)
    
    event = get_user_event(request.user)
    if not event:
        return HttpResponse('No hay evento configurado', status=400)
    
    wb = Workbook()
    ws = wb.active
    ws.title = 'Invitados'
    
    headers = ['nombre', 'email', 'telefono', 'grupo', 'estado', 'menu', 'acompañante', 'nombre_acompañante', 'notas']
    ws.append(headers)
    
    for guest in event.guests.filter(linked_to__isnull=True):
        ws.append([
            guest.full_name,
            guest.email,
            guest.phone,
            guest.get_group_display(),
            guest.get_rsvp_status_display(),
            guest.get_meal_choice_display(),
            'Sí' if guest.plus_one else 'No',
            guest.plus_one_name,
            guest.notes,
        ])
    
    for col in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 20
    
    ws.column_dimensions['A'].width = 25
    ws.column_dimensions['B'].width = 30
    ws.column_dimensions['I'].width = 30
    
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename=invitados_{event.slug}.xlsx'
    wb.save(response)
    return response


@login_required
def download_guest_template(request):
    if request.user.role != 'PLANNER':
        return HttpResponse('No autorizado', status=403)
    
    wb = Workbook()
    ws = wb.active
    ws.title = 'Invitados'
    
    headers = ['nombre', 'email', 'telefono', 'grupo', 'menu', 'acompañante', 'nombre_acompañante', 'notas']
    ws.append(headers)
    
    sample_data = [
        ['María García', 'maria@ejemplo.com', '555-1234', 'Familia Novia', 'Estándar', 'No', '', 'Alergia a mariscos'],
        ['Juan López', 'juan@ejemplo.com', '555-5678', 'Familia Novio', 'Vegetariano', 'Sí', 'Ana López', ''],
        ['Carlos Ruiz', 'carlos@ejemplo.com', '', 'Amigos Novia', 'Estándar', 'No', '', ''],
        ['Sofia Torres', 'sofia@ejemplo.com', '', 'Amigos Novio', 'Vegano', 'No', '', ''],
    ]
    
    for row in sample_data:
        ws.append(row)
    
    ws.append([])
    ws.append(['Grupos válidos:'])
    ws.append(['Familia Novia, Familia Novio, Amigos Novia, Amigos Novio, Trabajo, Otros'])
    ws.append([])
    ws.append(['Menús válidos:'])
    ws.append(['Estándar, Vegetariano, Vegano, Sin Gluten'])
    
    for col in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 20
    
    ws.column_dimensions['A'].width = 25
    ws.column_dimensions['B'].width = 30
    ws.column_dimensions['H'].width = 30
    
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename=plantilla_invitados.xlsx'
    wb.save(response)
    return response


class BudgetItemForm(forms.ModelForm):
    class Meta:
        model = BudgetItem
        fields = ['category', 'description', 'estimated_cost', 'actual_cost']
        widgets = {
            'description': forms.TextInput(attrs={'placeholder': 'Descripción del gasto'}),
            'estimated_cost': forms.NumberInput(attrs={'placeholder': '0', 'step': '0.01'}),
            'actual_cost': forms.NumberInput(attrs={'placeholder': '0', 'step': '0.01'}),
        }


class BudgetItemCreateForm(forms.ModelForm):
    class Meta:
        model = BudgetItem
        fields = ['category', 'description', 'estimated_cost']
        widgets = {
            'description': forms.TextInput(attrs={'placeholder': 'Descripción del gasto'}),
            'estimated_cost': forms.NumberInput(attrs={'placeholder': '0', 'step': '0.01'}),
        }


class BudgetView(PlannerRequiredMixin, ListView):
    template_name = 'dashboard/budget.html'
    model = BudgetItem
    context_object_name = 'budget_items'

    def get_queryset(self):
        event = get_user_event(self.request.user)
        if event:
            return event.budget_items.all()
        return BudgetItem.objects.none()

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        event = get_user_event(self.request.user)
        context['event'] = event
        if event:
            items = event.budget_items.all()
            context['total_estimated'] = sum(item.estimated_cost for item in items)
            context['total_actual'] = sum(item.actual_cost for item in items)
            context['budget_total'] = event.budget_total
        else:
            context['total_estimated'] = 0
            context['total_actual'] = 0
            context['budget_total'] = 0
        return context


class BudgetItemCreateView(PlannerRequiredMixin, CreateView):
    model = BudgetItem
    form_class = BudgetItemCreateForm
    template_name = 'dashboard/budget_form.html'

    def get_success_url(self):
        return reverse_lazy('dashboard:budget')

    def form_valid(self, form):
        event = get_user_event(self.request.user)
        if not event:
            messages.error(self.request, 'Primero debes configurar tu evento.')
            return redirect('dashboard:event_settings')
        form.instance.event = event
        messages.success(self.request, f'Item "{form.instance.description}" agregado al presupuesto.')
        return super().form_valid(form)

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['event'] = get_user_event(self.request.user)
        return context


class BudgetItemUpdateView(PlannerRequiredMixin, UpdateView):
    model = BudgetItem
    form_class = BudgetItemForm
    template_name = 'dashboard/budget_form.html'

    def get_success_url(self):
        return reverse_lazy('dashboard:budget')

    def form_valid(self, form):
        messages.success(self.request, f'Item "{form.instance.description}" actualizado.')
        return super().form_valid(form)

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['event'] = get_user_event(self.request.user)
        return context


class BudgetItemDeleteView(PlannerRequiredMixin, DeleteView):
    model = BudgetItem
    template_name = 'dashboard/budget_confirm_delete.html'

    def get_success_url(self):
        return reverse_lazy('dashboard:budget')

    def delete(self, request, *args, **kwargs):
        messages.success(request, 'Item de presupuesto eliminado.')
        return super().delete(request, *args, **kwargs)

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['event'] = get_user_event(self.request.user)
        return context


@login_required
def email_preview(request):
    if request.user.role != 'PLANNER':
        return redirect('dashboard:home')
    
    event = get_user_event(request.user)
    if not event:
        messages.error(request, 'Primero debes configurar tu evento.')
        return redirect('dashboard:event_settings')
    
    template_map = {
        'classic_romance': 'emails/invitations/classic_romance.html',
        'modern_minimal': 'emails/invitations/modern_minimal.html',
        'rustic_boho': 'emails/invitations/rustic_boho.html',
        'art_deco': 'emails/invitations/art_deco.html',
        'floral_romantic': 'emails/invitations/floral_romantic.html',
    }
    
    design = event.design_template or 'classic_romance'
    email_template = template_map.get(design, 'emails/invitations/classic_romance.html')
    
    preview_guest = Guest(
        full_name='Juan Pérez',
        email='juan@ejemplo.com',
        event=event,
    )
    
    from django.template.loader import render_to_string
    context = {
        'guest': preview_guest,
        'event': event,
        'invitation_url': 'https://ejemplo.com/invitacion/preview-token/',
    }
    
    html_content = render_to_string(email_template, context)
    return HttpResponse(html_content)
